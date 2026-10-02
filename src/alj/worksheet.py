"""Ergonomics of the manual annotation: read the filled worksheet, check it, measure progress, plan batches.

Nothing in this module produces, suggests or corrects a label (CLAUDE.md §2). It only:

* reads the worksheet the way the researcher keeps it: the ``gold_v1.xlsx`` he works in, or a CSV as Excel saves
  it (UTF-8 with or without BOM; ``,`` or ``;`` as separator — the pt-BR Excel writes ``;``);
* checks what the researcher typed against the codes of ``docs/annotation_protocol.md`` §§2–6 and lists
  **errors** (a code the protocol does not have, a missing column, a duplicated ``doc_id``) and **warnings**
  (combinations the protocol makes unlikely, to be looked at again — never changed automatically);
* counts done/remaining per stratum, in the order recommended by ``docs/COMO_ANOTAR.md``, with an estimate of
  the hours left at 2–4 minutes per document (and at the researcher's own measured pace, when ``minutes_spent``
  is filled);
* plans daily batches in that order, with a deterministic pseudo-random order *inside* each stratum, so that any
  prefix of a stratum is a random subsample of it (partial validations then stay unbiased; the worksheet itself
  is sorted by publication date, so its prefix would be the oldest decisions);
* dates the blind re-annotation round (protocol §6.3: at least two weeks after the first annotation of the same
  documents).
"""

from __future__ import annotations

import csv
import datetime as dt
import hashlib
import io
import re
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

from alj.annotation import DEFAULT_STRATA, LABEL_COLUMNS
from alj.validation import DOMAINS, MEASURES, SIGNALLING_STATUSES

#: order of work of docs/COMO_ANOTAR.md §4 (calibrate on the named phenomenon first, controls last)
RECOMMENDED_ORDER: tuple[str, ...] = (
    "strict",
    "strict_negated",
    "conduct_multi",
    "conduct_sanction",
    "control_flagged",
    "control_unflagged",
)
REQUIRED_COLUMNS: tuple[str, ...] = ("doc_id", "stratum", *LABEL_COLUMNS)
REANNOTATION_MIN_DAYS = 14  # protocol §6.3: "after two weeks"
MINUTES_PER_DOC = (2, 4)  # docs/COMO_ANOTAR.md §4, once the scale is calibrated
PROTOCOL_VERSIONS = frozenset({"v0.1"})

_GROUND_TOKEN = re.compile(r"^(A(?:[1-9]|1[0-9]|20)|T1198|MAFE)$")
_ANNEX_A = re.compile(r"^A(?:[1-9]|1[0-9]|20)$")

# ---------------------------------------------------------------------------------------------------- reading


def _cell(value) -> str:
    """Excel cells back to the text the researcher typed: dates as ISO, whole numbers without '.0'."""
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        return value.date().isoformat() if value.time() == dt.time() else value.isoformat(sep=" ")
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def read_worksheet(path: Path | str) -> tuple[list[str], list[dict]]:
    """Header and rows of the worksheet: a ``.xlsx`` (first sheet) or a CSV with ``,`` or ``;`` as separator."""
    path = Path(path)
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            it = wb.worksheets[0].iter_rows(values_only=True)
            header = [_cell(h).strip() for h in next(it, ())]
            rows = []
            for values in it:
                if values is None or all(v is None or str(v).strip() == "" for v in values):
                    continue
                rows.append({h: _cell(v) for h, v in zip(header, values, strict=False) if h})
        finally:
            wb.close()
        return [h for h in header if h], rows
    text = path.read_bytes().decode("utf-8-sig")
    first = text.split("\n", 1)[0]
    delimiter = ";" if first.count(";") > first.count(",") else ","
    reader = csv.DictReader(io.StringIO(text, newline=""), delimiter=delimiter)
    header = [h.strip() for h in (reader.fieldnames or [])]
    reader.fieldnames = header
    rows = [{k: (v if v is not None else "") for k, v in r.items() if k is not None} for r in reader]
    return header, rows


def parse_date(value: str | None) -> dt.date | None:
    """``2026-10-02``, ``02/10/2026`` (Excel pt-BR) or ``2026-10-02 14:30``; anything else → ``None``."""
    v = (value or "").strip()
    if not v:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d/%m/%y", "%Y/%m/%d"):
        try:
            return dt.datetime.strptime(v[:10] if fmt.startswith("%Y") else v.split()[0], fmt).date()
        except ValueError:
            continue
    return None


def is_labelled(row: dict) -> bool:
    """A row counts as done when ``label1_status`` is filled (blank = not annotated, COMO_ANOTAR §5)."""
    return bool((row.get("label1_status") or "").strip())


# --------------------------------------------------------------------------------------------------- checking


@dataclass(frozen=True)
class Issue:
    severity: str  # "erro" | "aviso"
    doc_id: str
    column: str
    message: str


def _split_codes(value: str) -> list[str]:
    return [p.strip().upper() for p in re.split(r"[;,/]", value) if p.strip()]


def check_header(header: list[str]) -> list[Issue]:
    missing = [c for c in REQUIRED_COLUMNS if c not in header]
    return [Issue("erro", "-", c, "coluna obrigatória ausente (o Excel pode ter renomeado ou apagado)")
            for c in missing]


def check_row(row: dict, *, known_strata: frozenset[str] | None = None) -> list[Issue]:
    """Errors and warnings for one row. Codes are compared with the protocol; nothing is rewritten."""
    doc = (row.get("doc_id") or "").strip() or "?"
    out: list[Issue] = []

    def err(col: str, msg: str) -> None:
        out.append(Issue("erro", doc, col, msg))

    def warn(col: str, msg: str) -> None:
        out.append(Issue("aviso", doc, col, msg))

    strata = known_strata if known_strata is not None else frozenset(s.name for s in DEFAULT_STRATA)
    stratum = (row.get("stratum") or "").strip()
    if stratum not in strata:
        warn("stratum", f"estrato desconhecido '{stratum}' (a coluna stratum não deve ser editada)")

    raw_status = (row.get("label1_status") or "").strip()
    status = raw_status.upper()
    grounds_raw = (row.get("label2_grounds") or "").strip()
    measure_raw = (row.get("label3_measure") or "").strip()
    domain_raw = (row.get("label4_domain") or "").strip()
    other_label_filled = any((row.get(c) or "").strip() for c in ("label2_grounds", "label3_measure",
                                                                   "label4_domain", "justification"))
    if not raw_status:
        if other_label_filled:
            warn("label1_status", "status em branco, mas outras colunas de rótulo preenchidas "
                                  "(a linha conta como NÃO anotada)")
        return out

    if status not in SIGNALLING_STATUSES:
        err("label1_status", f"'{raw_status}' não é código do protocolo §2 ({'/'.join(SIGNALLING_STATUSES)})")
        return out
    if raw_status != status:
        warn("label1_status", f"'{raw_status}' em minúsculas; use '{status}'")

    grounds = _split_codes(grounds_raw)
    bad = [g for g in grounds if not _GROUND_TOKEN.match(g)]
    if bad:
        err("label2_grounds", f"código(s) fora do §3: {', '.join(bad)} (use A1…A20, T1198, MAFE separados por ;)")
    dup = [g for g, n in Counter(grounds).items() if n > 1]
    if dup:
        warn("label2_grounds", f"código repetido: {', '.join(dup)}")
    annex = [g for g in grounds if _ANNEX_A.match(g)]

    measure = measure_raw.upper()
    if measure and measure not in MEASURES:
        err("label3_measure", f"'{measure_raw}' não é código do §4 ({'/'.join(MEASURES)}; escolha única)")
    domain = domain_raw.upper()
    if domain and domain not in DOMAINS:
        err("label4_domain", f"'{domain_raw}' não é código do §5 ({'/'.join(DOMAINS)}; escolha única)")

    minutes = (row.get("minutes_spent") or "").strip().replace(",", ".")
    if minutes:
        try:
            m = float(minutes)
            if m <= 0 or m > 240:
                warn("minutes_spent", f"{minutes} min parece fora do plausível")
        except ValueError:
            err("minutes_spent", f"'{row.get('minutes_spent')}' não é número de minutos")
    else:
        warn("minutes_spent", "vazio (§6.1 pede o tempo gasto)")

    if not (row.get("annotator") or "").strip():
        warn("annotator", "vazio (§6.4)")
    date_raw = (row.get("annotation_date") or "").strip()
    if not date_raw:
        warn("annotation_date", "vazio (§6.4; necessário para datar a rodada cega)")
    elif parse_date(date_raw) is None:
        err("annotation_date", f"'{date_raw}' não é data (use AAAA-MM-DD ou DD/MM/AAAA)")
    pv = (row.get("protocol_version") or "").strip()
    if not pv:
        warn("protocol_version", "vazio (§6.4; a versão confirmada é v0.1)")
    elif pv not in PROTOCOL_VERSIONS:
        warn("protocol_version", f"'{pv}' difere das versões confirmadas ({', '.join(sorted(PROTOCOL_VERSIONS))})")

    if status == "NA":
        if grounds or measure or domain:
            warn("label1_status", "NA (texto ausente/ilegível) com fundamentos, medida ou domínio preenchidos")
        return out

    if not (row.get("justification") or "").strip():
        warn("justification", "vazia (§6.2 pede uma frase com a passagem decisiva, por parágrafo)")
    if not domain:
        warn("label4_domain", "vazio (§5 é obrigatório para S3/S2/S1/S0)")
    if not measure:
        warn("label3_measure", "vazio (§4: use M0 quando não houver medida)")
    if status == "S3" and measure == "M0":
        warn("label3_measure", "S3 pressupõe medida fundada nos indícios (§2), mas a medida é M0")
    if status in {"S3", "S2"} and not grounds:
        warn("label2_grounds", f"{status} sem nenhum fundamento do Anexo A/T1198/MAFE (§3)")
    if status == "S0" and grounds:
        warn("label2_grounds", "S0 (falso positivo do léxico) com fundamentos preenchidos")
    if "MAFE" in grounds and annex:
        warn("label2_grounds", "MAFE é 'só má-fé, sem padrão de massificação' (§3), mas há itens A preenchidos")
    if measure == "M5" and not (row.get("justification") or "").strip():
        warn("label3_measure", "M5 (outra) pede o texto livre na justificativa (§4)")
    return out


def check_worksheet(header: list[str], rows: list[dict]) -> list[Issue]:
    issues = check_header(header)
    if any(i.severity == "erro" and i.column in {"doc_id", "label1_status"} for i in issues):
        return issues
    ids = Counter((r.get("doc_id") or "").strip() for r in rows)
    for doc, n in ids.items():
        if not doc:
            issues.append(Issue("erro", "?", "doc_id", f"{n} linha(s) sem doc_id"))
        elif n > 1:
            issues.append(Issue("erro", doc, "doc_id", f"doc_id repetido {n} vezes"))
    for r in rows:
        issues.extend(check_row(r))
    return issues


# --------------------------------------------------------------------------------------------------- progress


@dataclass
class StratumProgress:
    stratum: str
    total: int
    done: int
    na: int = 0
    minutes_logged: float = 0.0
    minutes_rows: int = 0

    @property
    def remaining(self) -> int:
        return self.total - self.done

    @property
    def pct(self) -> float:
        return round(100 * self.done / self.total, 1) if self.total else 0.0


@dataclass
class Progress:
    strata: list[StratumProgress] = field(default_factory=list)

    @property
    def total(self) -> int:
        return sum(s.total for s in self.strata)

    @property
    def done(self) -> int:
        return sum(s.done for s in self.strata)

    @property
    def remaining(self) -> int:
        return self.total - self.done

    def own_pace(self) -> float | None:
        """Mean ``minutes_spent`` over the rows where it was filled (None when nothing was logged yet)."""
        n = sum(s.minutes_rows for s in self.strata)
        return round(sum(s.minutes_logged for s in self.strata) / n, 2) if n else None

    def hours_left(self, minutes_per_doc: float) -> float:
        return round(self.remaining * minutes_per_doc / 60, 1)

    def as_dict(self) -> dict:
        lo, hi = MINUTES_PER_DOC
        pace = self.own_pace()
        return {
            "total": self.total,
            "done": self.done,
            "remaining": self.remaining,
            "pct_done": round(100 * self.done / self.total, 1) if self.total else 0.0,
            "hours_left_at_2_to_4_min": [self.hours_left(lo), self.hours_left(hi)],
            "own_pace_min_per_doc": pace,
            "hours_left_at_own_pace": self.hours_left(pace) if pace else None,
            "by_stratum": [
                {"stratum": s.stratum, "total": s.total, "done": s.done, "remaining": s.remaining,
                 "pct": s.pct, "na": s.na} for s in self.strata
            ],
        }


def stratum_rank(name: str) -> int:
    return RECOMMENDED_ORDER.index(name) if name in RECOMMENDED_ORDER else len(RECOMMENDED_ORDER)


def progress(rows: list[dict]) -> Progress:
    by: dict[str, StratumProgress] = {}
    for r in rows:
        name = (r.get("stratum") or "").strip() or "?"
        sp = by.setdefault(name, StratumProgress(name, 0, 0))
        sp.total += 1
        if is_labelled(r):
            sp.done += 1
            sp.na += (r.get("label1_status") or "").strip().upper() == "NA"
            try:
                m = float((r.get("minutes_spent") or "").strip().replace(",", "."))
                if 0 < m <= 240:
                    sp.minutes_logged += m
                    sp.minutes_rows += 1
            except ValueError:
                pass
    return Progress(sorted(by.values(), key=lambda s: (stratum_rank(s.stratum), s.stratum)))


# ---------------------------------------------------------------------------------------------------- batches


def order_key(doc_id: str, seed: int) -> str:
    """Deterministic pseudo-random position inside a stratum (same seed → same order on every machine)."""
    return hashlib.sha256(f"{seed}:{doc_id}".encode()).hexdigest()


def plan_batches(rows: list[dict], per_day: int = 25, seed: int = 20260922) -> list[dict]:
    """One entry per document: ``ordem`` (global position), ``lote`` (1-based), ``stratum`` and ``doc_id``.

    Strata follow :data:`RECOMMENDED_ORDER`; inside a stratum the order is pseudo-random (``order_key``), so
    stopping halfway through a stratum still leaves a random subsample of it. Rows already labelled keep their
    place in the plan (the plan does not depend on the labels) and are reported as done.
    """
    if per_day < 1:
        raise ValueError("per_day must be >= 1")
    ordered = sorted(
        rows,
        key=lambda r: (stratum_rank((r.get("stratum") or "").strip()), (r.get("stratum") or "").strip(),
                       order_key((r.get("doc_id") or "").strip(), seed)),
    )
    return [
        {"ordem": i + 1, "lote": i // per_day + 1, "stratum": (r.get("stratum") or "").strip(),
         "doc_id": (r.get("doc_id") or "").strip(), "feito": is_labelled(r)}
        for i, r in enumerate(ordered)
    ]


def batch_dates(n_batches: int, start: dt.date, weekdays_only: bool = True) -> list[dt.date]:
    out: list[dt.date] = []
    day = start
    while len(out) < n_batches:
        if not weekdays_only or day.weekday() < 5:
            out.append(day)
        day += dt.timedelta(days=1)
    return out


# ------------------------------------------------------------------------------------------ blind re-annotation


def reannotation_schedule(gold_rows: list[dict], reann_ids: list[str],
                          min_days: int = REANNOTATION_MIN_DAYS) -> dict:
    """When the blind round (§6.3) may start: ≥ ``min_days`` after the LAST first-round date of those documents.

    Reports how many of the re-annotation documents are already done in the first round, how many lack a
    readable ``annotation_date`` and the earliest allowed start date (``None`` until all are done and dated).
    """
    by_id = {(r.get("doc_id") or "").strip(): r for r in gold_rows}
    ids = [i.strip() for i in reann_ids if i.strip()]
    done = [by_id[i] for i in ids if i in by_id and is_labelled(by_id[i])]
    dates = [parse_date(r.get("annotation_date")) for r in done]
    undated = sum(d is None for d in dates)
    known = [d for d in dates if d is not None]
    pending = len(ids) - len(done)
    earliest = (max(known) + dt.timedelta(days=min_days)) if known and not pending and not undated else None
    return {
        "reannotation_docs": len(ids),
        "missing_from_gold": sum(i not in by_id for i in ids),
        "done_in_first_round": len(done),
        "pending_in_first_round": pending,
        "done_without_date": undated,
        "last_first_round_date": max(known).isoformat() if known else None,
        "earliest_blind_start": earliest.isoformat() if earliest else None,
        "min_days": min_days,
    }


def reannotation_gaps(gold_rows: list[dict], second_rows: list[dict],
                      min_days: int = REANNOTATION_MIN_DAYS) -> dict:
    """For a (partially) filled blind round: pairs re-annotated too soon after the first round (< ``min_days``)."""
    first = {(r.get("doc_id") or "").strip(): parse_date(r.get("annotation_date")) for r in gold_rows}
    too_soon, undated, pairs = 0, 0, 0
    for r in second_rows:
        if not is_labelled(r):
            continue
        doc = (r.get("doc_id") or "").strip()
        d1, d2 = first.get(doc), parse_date(r.get("annotation_date"))
        pairs += 1
        if d1 is None or d2 is None:
            undated += 1
        elif (d2 - d1).days < min_days:
            too_soon += 1
    return {"pairs": pairs, "too_soon": too_soon, "undated": undated, "min_days": min_days}
